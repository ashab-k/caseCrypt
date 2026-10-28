"""Generate fictional upload fixtures. Run with Python + python-docx."""
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1] / 'docs' / 'demo'
DOCUMENTS = [
    ('01_mock_complaint', 'Complaint and statement of facts', 'DEMO/2023/017', [
        ('Case particulars', [
            'Petitioner: Asha Mehta', 'Respondent: Rohan Shah', 'Counsel: Priya Rao',
            'Forum: Fictional District Court, Demo City',
        ]),
        ('Statement of facts', [
            'On 12 June 2023, Asha Mehta transferred INR 85,000 to Rohan Shah for the supply of five office laptops.',
            'On 19.06.2023, the alleged non-delivery incident occurred at the agreed delivery address in Demo City.',
            'On 22/06/2023, Asha Mehta filed a written complaint with the Demo City Police Station.',
        ]),
        ('Claim and references', [
            'The petitioner alleges that the respondent accepted payment but did not deliver the laptops or refund INR 85,000.',
            'The complaint cites Section 420 of the Indian Penal Code and Section 154 of the Code of Criminal Procedure.',
            'These references are fictional pleadings for extraction testing, not findings of an offence.',
        ]),
        ('Documents referred to', [
            'Exhibit A: Mock payment receipt for INR 85,000.',
            'Exhibit B: Mock purchase correspondence concerning five office laptops.',
        ]),
    ]),
    ('02_mock_hearing_order', 'Hearing record and procedural order', 'DEMO/2023/017', [
        ('Case particulars', [
            'Petitioner: Asha Mehta', 'Respondent: Rohan Shah', 'Counsel: Priya Rao',
            'Forum: Fictional District Court, Demo City',
        ]),
        ('Procedural record', [
            'On 03 July 2023, the court held a preliminary hearing with both parties present.',
            'On 10.07.2023, the court ordered both parties to produce the payment receipt and purchase correspondence.',
            'On 24-07-2023, the court held a further hearing and recorded receipt of the requested documents.',
        ]),
        ('Submissions recorded', [
            'The respondent submits that a supplier delay prevented delivery and disputes any intention to deceive.',
            'The petitioner claims a refund of INR 85,000 and relies on the payment receipt.',
            'The petitioner refers again to Section 420 IPC. No decision on liability is recorded in this fictional procedural order.',
        ]),
        ('Scope of the record', [
            'The record contains procedural steps only. The allegations remain disputed.',
            'All names, events, exhibits, payments, and court details in this record are invented for the CaseCrypt demonstration.',
        ]),
    ]),
]


def main():
    ROOT.mkdir(parents=True, exist_ok=True)
    for stem, title, reference, sections in DOCUMENTS:
        doc = Document()
        section = doc.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.65)
        section.left_margin = section.right_margin = Inches(0.8)
        for style in doc.styles:
            for border in list(style.element.iter(qn('w:pBdr'))):
                border.getparent().remove(border)
        normal = doc.styles['Normal']
        normal.font.name = 'Calibri'
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.08
        for name in ('Title', 'Subtitle', 'Heading 1', 'Heading 2'):
            doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles['Title'].font.size = Pt(23)
        doc.styles['Heading 1'].font.size = Pt(12)
        doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
        doc.styles['Heading 1'].paragraph_format.space_after = Pt(5)
        notice = doc.add_paragraph('FICTIONAL CASE • DEMONSTRATION ONLY')
        notice.runs[0].bold = True
        notice.runs[0].font.size = Pt(9)
        doc.add_paragraph(title, 'Title')
        doc.add_paragraph(f'CaseCrypt demo bundle | Case reference: {reference}', 'Subtitle')
        for heading, paragraphs in sections:
            doc.add_heading(heading, level=1)
            for text in paragraphs:
                doc.add_paragraph(text)
        doc.core_properties.author = 'CaseCrypt'
        doc.core_properties.title = title
        doc.core_properties.subject = 'Fictional demonstration case; no real people or proceedings'
        doc.save(ROOT / f'{stem}.docx')
        print(ROOT / f'{stem}.docx')


if __name__ == '__main__':
    main()
