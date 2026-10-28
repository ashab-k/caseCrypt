"""Exercise actual DOCX ingestion and the offline demo extraction together."""
import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

from docx import Document

from services.nlp_service import _extract_heuristic
from services.ocr_service import extract_text, pages_to_raw_text
from services.timeline_service import aggregate_entities, build_timeline

DEMO = Path(__file__).resolve().parents[2] / 'docs' / 'demo'


class DemoDocumentTests(unittest.TestCase):
    def test_demo_bundle_has_six_dated_events_and_two_provisions(self):
        pairs = []
        for path in sorted(DEMO.glob('*.docx')):
            payload = _extract_heuristic(pages_to_raw_text(extract_text(path, 'docx')), [])
            extraction = SimpleNamespace(extracted_at=datetime(2023, 7, 25), extraction_method='heuristic',
                **{f'{key}_json': json.dumps(payload.get(key, [])) for key in
                   ['events', 'dates', 'parties', 'claims', 'cited_provisions', 'spacy_entities']})
            document = SimpleNamespace(id=path.stem, original_name=path.name, file_type='docx')
            pairs.append((extraction, document))
        self.assertEqual(len(pairs), 2)
        timeline = build_timeline(pairs)
        self.assertEqual([e['date'] for e in timeline['events']],
                         ['2023-06-12', '2023-06-19', '2023-06-22', '2023-07-03', '2023-07-10', '2023-07-24'])
        entities = aggregate_entities(pairs)
        self.assertEqual(entities['parties'], {'petitioner': ['Asha Mehta'], 'respondent': ['Rohan Shah'], 'counsel': ['Priya Rao']})
        self.assertEqual([(p['act'], p['section']) for p in entities['cited_provisions']], [('IPC', '420'), ('CrPC', '154')])
        self.assertEqual(len(entities['claims']), 3)

    def test_docx_tables_are_extracted_in_order(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'chronology.docx'
            doc = Document()
            doc.add_paragraph('Petitioner: Asha Mehta')
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = '25.08.2020'
            table.cell(0, 1).text = 'The petitioner filed a complaint.'
            doc.add_paragraph('After the table.')
            doc.save(path)
            text = pages_to_raw_text(extract_text(path, 'docx'))
            self.assertLess(text.index('Asha Mehta'), text.index('25.08.2020'))
            self.assertLess(text.index('filed a complaint'), text.index('After the table'))
            self.assertEqual(_extract_heuristic(text, [])['events'][0]['date'], '2020-08-25')
